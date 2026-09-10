"""Read all original XLSX inputs and emit reproducible schema/hash evidence."""

import hashlib
import json
import math
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ATTACHMENTS = ROOT / "附件"


def require(condition, message):
    if not condition:
        raise ValueError(message)


def summarize(values):
    require(all(isinstance(v, (int, float)) and not isinstance(v, bool)
                and math.isfinite(v) for v in values), "missing or invalid numeric value")
    require(min(values) >= 0, "negative input")
    return {"count": len(values), "min": min(values), "max": max(values)}


def expected_labels():
    return [(datetime(2025, 1, 1) + timedelta(minutes=10 * t)).time()
            for t in range(1, 144)] + ["0:00+1"]


def normalize_time_label(value):
    if value == "0:00+1":
        return value
    if isinstance(value, str):
        return datetime.strptime(value.strip(), "%H:%M").time()
    return value


def dates_2025():
    return [datetime(2025, 1, 1) + timedelta(days=d) for d in range(365)]


def main():
    report = {"kind": "original_input_audit_not_optimization", "cycle_records": 144,
              "cycle_first_label": "0:10", "cycle_last_label": "0:00+1",
              "evaluation_days": 334, "source_hashes": {}, "sheets": {},
              "template_differences": []}
    files = [ROOT / "C题.pdf", *sorted(p for p in ATTACHMENTS.rglob("*.xlsx") if not p.name.startswith("~$"))]
    for path in files:
        report["source_hashes"][path.relative_to(ROOT).as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
    for path in sorted(p for p in ATTACHMENTS.glob("*.xlsx") if not p.name.startswith("~$")):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = list(sheet.values)
                key = f"{path.name}/{sheet.title}"
                item = {"rows_including_header": len(rows), "columns": len(rows[0])}
                if path.name == "附件1.xlsx":
                    require(len(rows) == 145 and len(rows[0]) == 4, "attachment 1 shape")
                    require([normalize_time_label(r[0]) for r in rows[1:]] == expected_labels(), "attachment 1 time labels")
                    item["time_storage_types"] = sorted({type(r[0]).__name__ for r in rows[1:]})
                    item["fields"] = {rows[0][j]: summarize([r[j] for r in rows[1:]]) for j in range(1, 4)}
                elif path.name in {"附件2.xlsx", "附件4.xlsx"}:
                    require(len(rows) == 366 and len(rows[0]) == 145, f"{key}: shape")
                    require([normalize_time_label(v) for v in rows[0][1:]] == expected_labels(), f"{key}: time labels")
                    require([r[0] for r in rows[1:]] == dates_2025(), f"{key}: dates")
                    item["values"] = summarize([v for r in rows[1:] for v in r[1:]])
                    item["date_start"] = str(rows[1][0].date())
                    item["date_end"] = str(rows[-1][0].date())
                elif path.name == "附件3.xlsx":
                    require(len(rows) == 1461 and len(rows[0]) == 26, "forecast shape")
                    require(list(rows[0][2:]) == [f"预报{h}小时" for h in range(1, 25)], "forecast lead headers")
                    current_date = None
                    issues = []
                    for row in rows[1:]:
                        if row[0] not in (None, ""):
                            current_date = datetime.strptime(str(row[0]), "%Y-%m-%d")
                        require(current_date is not None, "forecast date unavailable")
                        release = datetime.strptime(str(row[1]), "%H:%M").time()
                        issues.append(datetime.combine(current_date.date(), release))
                    expected = [d + timedelta(hours=h) for d in dates_2025() for h in (0, 6, 12, 18)]
                    require(issues == expected, "forecast releases missing, duplicated or reordered")
                    item["values"] = summarize([v for row in rows[1:] for v in row[2:]])
                    item["release_count"] = len(issues)
                    item["last_valid_time"] = (issues[-1] + timedelta(hours=24)).isoformat()
                report["sheets"][key] = item
        finally:
            workbook.close()

    expected_sheets = {"计划购电量", "充放电量", "紧急购电量"}
    for path in sorted(p for p in (ATTACHMENTS / "附件5").glob("*.xlsx") if not p.name.startswith("~$")):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            required = ({"计划购电量", "充放电量"} if path.name == "result1.xlsx" else
                        expected_sheets | ({"调整购电量"} if path.name in {"result3.xlsx", "result4-3.xlsx"} else set()))
            require(set(workbook.sheetnames) == required, f"{path.name}: sheet names")
            for sheet in workbook:
                rows = list(sheet.values)
                key = f"附件5/{path.name}/{sheet.title}"
                item = {"rows_including_header": len(rows), "columns": len(rows[0])}
                if sheet.title in {"计划购电量", "调整购电量"}:
                    if path.name == "result1.xlsx":
                        require(len(rows) == 145, "result1 records")
                        labels = [r[0] for r in rows[1:]]
                    else:
                        require(len(rows) == 335 and len(rows[0]) == 147, "annual template shape")
                        require([r[0] for r in rows[1:]] == dates_2025()[31:], "template evaluation dates")
                        labels = list(rows[0][1:145])
                        item["summary_headers"] = list(rows[0][145:147])
                    item["first_interval_label"] = labels[0]
                    item["last_interval_label"] = labels[-1]
                    report["template_differences"].append({"sheet": key, "first": labels[0], "last": labels[-1],
                        "resolution": "Use user-confirmed 144 endpoint records; preserve original template; explicitly map output labels."})
                item["contains_ellipsis"] = any(v == "⁝" for r in rows for v in r)
                report["sheets"][key] = item
        finally:
            workbook.close()
    require(report["sheets"]["附件4.xlsx/Sheet1"]["values"]["min"] > 0, "nonpositive price")
    report["status"] = "passed_with_documented_template_label_differences"
    target = ROOT / "docs" / "input_audit.json"
    target.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "hashed_source_files": len(files),
                      "inspected_sheets": len(report["sheets"]), "report": "docs/input_audit.json"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
