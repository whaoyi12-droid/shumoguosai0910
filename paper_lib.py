# -*- coding: utf-8 -*-
"""论文排版原语: 按国赛范文样式(A4, 2.5cm, 宋体/黑体, 三线表, 页码)。"""
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def new_doc():
    doc=Document()
    sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
    sec.header_distance=Cm(1.5);sec.footer_distance=Cm(1.5)
    st=doc.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(12)
    st._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    return doc

def _set(run,cn='宋体',en='Times New Roman',size=12,bold=False,italic=False,color=None):
    run.font.name=en;run.font.size=Pt(size);run.bold=bold;run.italic=italic
    run._element.rPr.rFonts.set(qn('w:eastAsia'),cn)
    if color:run.font.color.rgb=RGBColor(*color)

def para(doc,text='',size=12,cn='宋体',bold=False,align='just',indent=True,
         before=0,after=0,line=1.5,italic=False):
    p=doc.add_paragraph();pf=p.paragraph_format
    pf.space_before=Pt(before);pf.space_after=Pt(after)
    pf.line_spacing_rule=WD_LINE_SPACING.MULTIPLE;pf.line_spacing=line
    am={'just':WD_ALIGN_PARAGRAPH.JUSTIFY,'center':WD_ALIGN_PARAGRAPH.CENTER,
        'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT}
    p.alignment=am[align]
    if indent:p.paragraph_format.first_line_indent=Pt(size*2)
    if text!='':
        # 支持 **加粗** 片段
        import re
        for seg,b in [(s,i%2==1) for i,s in enumerate(re.split(r'\*\*',text))]:
            r=p.add_run(seg);_set(r,cn=cn,size=size,bold=bold or b,italic=italic)
    return p

def rich(doc,parts,size=12,align='just',indent=True,line=1.5,after=0):
    """parts: [(text,bold,italic)]"""
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if align=='just' else WD_ALIGN_PARAGRAPH.CENTER
    if indent:p.paragraph_format.first_line_indent=Pt(size*2)
    p.paragraph_format.line_spacing=line;p.paragraph_format.space_after=Pt(after)
    for t in parts:
        txt=t[0];b=t[1] if len(t)>1 else False;it=t[2] if len(t)>2 else False
        cn=t[3] if len(t)>3 else '宋体'
        r=p.add_run(txt);_set(r,cn=cn,size=size,bold=b,italic=it)
    return p

def h1(doc,text):
    p=para(doc,text,size=15,cn='黑体',bold=True,align='center',indent=False,before=10,after=8,line=1.5)
    return p
def h2(doc,text):
    return para(doc,text,size=13,cn='黑体',bold=True,align='left',indent=False,before=6,after=4)
def h3(doc,text):
    return para(doc,text,size=12,cn='黑体',bold=True,align='left',indent=False,before=4,after=2)

def formula(doc,text,num=None):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
    r=p.add_run(text);_set(r,en='Times New Roman',cn='宋体',size=12,italic=False)
    if num:
        rr=p.add_run('\t\t('+str(num)+')');_set(rr,size=12)
    return p

def caption(doc,text):
    return para(doc,text,size=10.5,cn='宋体',bold=True,align='center',indent=False,before=3,after=6)

def _set_cell_border(cell,**kw):
    tc=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
    for edge in ('top','bottom'):
        if edge in kw:
            e=OxmlElement('w:'+edge);a=kw[edge]
            for k,v in (('w:val','single'),('w:sz',str(a[0])),('w:space','0'),('w:color','000000')):e.set(qn(k),v)
            bd.append(e)
    tc.append(bd)

def three_table(doc,headers,rows,widths=None,fontsize=10.5,caption_text=None):
    if caption_text:caption(doc,caption_text)
    t=doc.add_table(rows=1+len(rows),cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=True
    def fill(cell,text,bold=False):
        if isinstance(text,float): text=('%.2f'%text).rstrip('0').rstrip('.') if False else '%.2f'%text
        cell.text='';p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_before=Pt(1);p.paragraph_format.space_after=Pt(1)
        r=p.add_run(str(text));_set(r,size=fontsize,bold=bold)
    for j,htxt in enumerate(headers):
        fill(t.rows[0].cells[j],htxt,bold=True)
        _set_cell_border(t.rows[0].cells[j],top=(12,),bottom=(6,))
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            fill(t.rows[i+1].cells[j],v)
            if i==len(rows)-1:_set_cell_border(t.rows[i+1].cells[j],bottom=(12,))
    if widths:
        for j,w in enumerate(widths):
            for r_ in t.rows:r_.cells[j].width=Cm(w)
    return t

def add_page_number(doc):
    sec=doc.sections[0];ftr=sec.footer;p=ftr.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=p.add_run();fld1=OxmlElement('w:fldChar');fld1.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text='PAGE'
    fld2=OxmlElement('w:fldChar');fld2.set(qn('w:fldCharType'),'end')
    run._r.append(fld1);run._r.append(instr);run._r.append(fld2);_set(run,size=10.5)

def header_rule(doc,text=''):
    sec=doc.sections[0];hdr=sec.header;p=hdr.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pPr=p._p.get_or_add_pPr();pbdr=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom')
    bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'6');bottom.set(qn('w:space'),'1'),bottom.set(qn('w:color'),'000000')
    pbdr.append(bottom);pPr.append(pbdr)

def code_block(doc,code,size=8):
    for line in code.split('\n'):
        p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0)
        r=p.add_run(line if line else ' ');r.font.name='Consolas';r.font.size=Pt(size)
        r._element.rPr.rFonts.set(qn('w:eastAsia'),'Consolas')
